from playwright.sync_api import sync_playwright
import time
import os
import logging
import json

logger = logging.getLogger(__name__)

class KDocsDownloader:
    def __init__(self):
        self.browser = None
        self.context = None
        self.page = None
        self.playwright = None
        self.debug_mode = True
    
    def _log(self, msg):
        if self.debug_mode:
            print(f"[KDocsDownloader] {msg}")
        logger.info(msg)
    
    def _setup_browser(self, headless=False):
        self._log("Setting up browser...")
        self.playwright = sync_playwright().start()
        self.browser = self.playwright.chromium.launch(
            headless=headless,
            args=["--start-maximized"],
            timeout=60000
        )
        
        download_path = os.path.join(os.path.dirname(__file__), "..", "downloads")
        os.makedirs(download_path, exist_ok=True)
        self._log(f"Download path: {download_path}")
        
        self.context = self.browser.new_context(
            viewport={"width": 1920, "height": 1080},
            accept_downloads=True,
            user_agent="Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/120.0.0.0 Safari/537.36"
        )
        
        storage_file = os.path.join(os.path.dirname(__file__), "..", "browser_storage.json")
        if os.path.exists(storage_file):
            try:
                with open(storage_file, "r", encoding="utf-8") as f:
                    storage_state = json.load(f)
                    self.context.storage_state(storage_state=storage_state)
                self._log("Loaded saved browser storage")
            except:
                self._log("Failed to load saved browser storage")
        else:
            self._log("No saved browser storage found")
        
        self.page = self.context.new_page()
        self.page.on("download", self._handle_download)
        self.downloaded_file = None
        self._log("Browser setup complete")
    
    def _handle_download(self, download):
        self._log(f"Download started: {download.url}")
        download_path = os.path.join(os.path.dirname(__file__), "..", "downloads")
        filename = download.suggested_filename or "download.xlsx"
        full_path = os.path.join(download_path, filename)
        download.save_as(full_path)
        self.downloaded_file = full_path
        logger.info(f"Download saved to: {full_path}")
    
    def _save_storage(self):
        try:
            storage_state = self.context.storage_state()
            storage_file = os.path.join(os.path.dirname(__file__), "..", "browser_storage.json")
            with open(storage_file, "w", encoding="utf-8") as f:
                json.dump(storage_state, f, ensure_ascii=False, indent=2)
            logger.info("Browser storage saved")
        except Exception as e:
            logger.warning(f"Failed to save storage: {e}")
    
    def _wait_for_login(self, timeout=120):
        self._log(f"Waiting for login, timeout: {timeout} seconds")
        start_time = time.time()
        
        while time.time() - start_time < timeout:
            try:
                body_text = self.page.inner_text("body")
                body_len = len(body_text)
                if body_len > 2000:
                    self._log(f"Login detected! Body text length: {body_len}")
                    self._save_storage()
                    return True
                
                current_url = self.page.url
                if "/doc/" in current_url or "/sheet/" in current_url:
                    self._log(f"URL indicates logged in: {current_url}")
                    return True
            except Exception as e:
                self._log(f"Error checking login: {e}")
            
            time.sleep(3)
            
            elapsed = int(time.time() - start_time)
            if elapsed % 30 == 0:
                self._log(f"Waiting... ({elapsed}s / {timeout}s)")
        
        self._log("Login timeout")
        return False
    
    def download_excel(self, document_url, sheet_name=None):
        self._log(f"Starting download from: {document_url}")
        
        try:
            self._setup_browser(headless=False)
            
            self._log(f"Opening URL: {document_url}")
            self.page.goto(document_url, wait_until="networkidle", timeout=120000)
            
            if not self._wait_for_login():
                self._log("Login failed")
                return None
            
            time.sleep(3)
            
            if sheet_name:
                self._log(f"Trying to select sheet: {sheet_name}")
                self._select_sheet(sheet_name)
                time.sleep(3)
            
            self._log("Looking for download button...")
            download_button_found = False
            
            selectors = [
                "button:has-text('下载')",
                "button:has-text('Download')",
                "[class*='download']",
                "[title*='下载']",
                ".wps-btn-download",
                "[data-action='download']",
                "a[download]",
                "[aria-label*='下载']"
            ]
            
            self._log(f"Trying {len(selectors)} selectors...")
            for i, selector in enumerate(selectors):
                try:
                    button = self.page.query_selector(selector)
                    if button:
                        self._log(f"✓ Found download button using selector [{i+1}]: {selector}")
                        button.click()
                        download_button_found = True
                        break
                except Exception as e:
                    self._log(f"✗ Selector [{i+1}] failed: {selector} - {e}")
            
            if not download_button_found:
                self._log("Trying to find download menu...")
                menu_selectors = [
                    "button:has-text('文件')",
                    "button:has-text('File')",
                    "[class*='menu']",
                    "[class*='toolbar']"
                ]
                
                for selector in menu_selectors:
                    try:
                        button = self.page.query_selector(selector)
                        if button:
                            self._log(f"Opening menu: {selector}")
                            button.click()
                            time.sleep(2)
                            
                            for dl_selector in selectors:
                                try:
                                    dl_button = self.page.query_selector(dl_selector)
                                    if dl_button:
                                        self._log(f"Found download in menu")
                                        dl_button.click()
                                        download_button_found = True
                                        break
                                except:
                                    pass
                            
                            if download_button_found:
                                break
                    except Exception as e:
                        self._log(f"Menu selector failed: {selector} - {e}")
            
            if not download_button_found:
                self._log("Trying keyboard shortcut Ctrl+S...")
                self.page.keyboard.press("Control+S")
                download_button_found = True
            
            if download_button_found:
                self._log("Waiting for download...")
                timeout = 120
                start_time = time.time()
                while time.time() - start_time < timeout:
                    if self.downloaded_file:
                        self._log(f"✓ Download completed: {self.downloaded_file}")
                        break
                    time.sleep(2)
                    elapsed = int(time.time() - start_time)
                    if elapsed % 15 == 0:
                        self._log(f"  Waiting... ({elapsed}s / {timeout}s)")
                
                if not self.downloaded_file:
                    self._log("Download may not have started, trying JS trigger...")
                    try:
                        js_download = """
                        var downloadLinks = document.querySelectorAll('a[download], [data-download]');
                        if (downloadLinks.length > 0) {
                            downloadLinks[0].click();
                            return 'Clicked download link';
                        }
                        return 'No download link found';
                        """
                        result = self.page.evaluate(js_download)
                        self._log(f"JS download result: {result}")
                        
                        time.sleep(10)
                        
                        if self.downloaded_file:
                            self._log(f"✓ Download completed after JS trigger: {self.downloaded_file}")
                        else:
                            self._log("✗ Download still not started")
                    except Exception as e:
                        self._log(f"JS trigger failed: {e}")
            
            if not self.downloaded_file:
                self._log("=" * 50)
                self._log("⚠ WARNING: Auto-download failed!")
                self._log("Please manually download the file:")
                self._log("1. Click 'File' or '文件' menu")
                self._log("2. Select 'Download' or '下载'")
                self._log("3. Choose Excel format (.xlsx)")
                self._log("4. Save to the 'downloads' folder")
                self._log("=" * 50)
                
                wait_time = 120
                self._log(f"Waiting {wait_time} seconds for manual download...")
                start_time = time.time()
                download_path = os.path.join(os.path.dirname(__file__), "..", "downloads")
                
                while time.time() - start_time < wait_time:
                    try:
                        files = os.listdir(download_path)
                        xlsx_files = [f for f in files if f.endswith('.xlsx') or f.endswith('.xls')]
                        if xlsx_files:
                            newest_file = max(xlsx_files, key=lambda f: os.path.getmtime(os.path.join(download_path, f)))
                            full_path = os.path.join(download_path, newest_file)
                            self.downloaded_file = full_path
                            self._log(f"✓ Found manually downloaded file: {newest_file}")
                            break
                    except:
                        pass
                    
                    time.sleep(3)
                    elapsed = int(time.time() - start_time)
                    if elapsed % 30 == 0:
                        self._log(f"  Waiting for manual download... ({elapsed}s / {wait_time}s)")
            
            self._save_storage()
            return self.downloaded_file
            
        except Exception as e:
            self._log(f"✗ Download failed: {e}")
            logger.error(f"Download failed: {e}", exc_info=True)
            return None
        finally:
            if self.playwright:
                try:
                    self.playwright.stop()
                except:
                    pass
    
    def _select_sheet(self, sheet_name):
        selectors = [
            f"div[title='{sheet_name}']",
            f"span[title='{sheet_name}']",
            f"button[title='{sheet_name}']",
            f"div[title*='{sheet_name}']"
        ]
        
        for selector in selectors:
            try:
                element = self.page.query_selector(selector)
                if element:
                    element.click()
                    self._log(f"Switched to sheet: {sheet_name}")
                    return True
            except Exception as e:
                pass
        
        self._log(f"Sheet '{sheet_name}' not found")
        return False
    
    def close(self):
        if self.page:
            try:
                self.page.close()
            except:
                pass
        if self.context:
            try:
                self.context.close()
            except:
                pass
        if self.browser:
            try:
                self.browser.close()
            except:
                pass
        if self.playwright:
            try:
                self.playwright.stop()
            except:
                pass

def extract_from_excel(file_path, sheet_name=None):
    try:
        import pandas as pd
        
        logger.info(f"Reading Excel file: {file_path}")
        
        if sheet_name:
            df = pd.read_excel(file_path, sheet_name=sheet_name)
        else:
            df = pd.read_excel(file_path)
        
        table_data = []
        headers = df.columns.tolist()
        table_data.append(headers)
        
        for _, row in df.iterrows():
            row_data = []
            for col in headers:
                val = row[col]
                if pd.isna(val):
                    row_data.append("")
                else:
                    row_data.append(str(val))
            table_data.append(row_data)
        
        logger.info(f"Extracted {len(table_data)} rows, {len(headers)} columns")
        return [table_data]
        
    except ImportError:
        logger.error("pandas not installed, trying openpyxl...")
        return extract_with_openpyxl(file_path, sheet_name)
    except Exception as e:
        logger.error(f"Error reading Excel: {e}")
        return None

def extract_with_openpyxl(file_path, sheet_name=None):
    try:
        from openpyxl import load_workbook
        
        logger.info(f"Reading Excel with openpyxl: {file_path}")
        wb = load_workbook(file_path, read_only=True)
        
        if sheet_name and sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
        else:
            ws = wb.active
        
        table_data = []
        for row in ws.iter_rows(values_only=True):
            row_data = []
            for cell in row:
                if cell is None:
                    row_data.append("")
                else:
                    row_data.append(str(cell))
            if any(row_data):
                table_data.append(row_data)
        
        logger.info(f"Extracted {len(table_data)} rows")
        return [table_data]
        
    except Exception as e:
        logger.error(f"Error reading Excel with openpyxl: {e}")
        return None

def download_and_extract(document_url, sheet_name=None):
    downloader = KDocsDownloader()
    file_path = downloader.download_excel(document_url, sheet_name)
    
    if not file_path:
        logger.error("Download failed")
        return None
    
    return extract_from_excel(file_path, sheet_name)