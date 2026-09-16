# -*- coding: utf-8 -*-
"""金山文档(kdocs)在线表格读取 —— Cookie 凭证方案。

不再走 developer.kdocs.cn OAuth2（个人开发者审核不通过），改为使用浏览器
登录态 wps_sid Cookie 调用金山文档网页版接口（与官方灵犀 SDK 同一机制）：

    浏览器 F12 → Application → Cookies → www.kdocs.cn → 复制 wps_sid
    → 程序带 Cookie 调 drive.kdocs.cn 接口 → 下载 xlsx 到 %TEMP%
    → 解析全部 sheet → 读完立即删除临时文件

注意：接口为网页版内部接口，若金山改版导致失败，日志会输出原始响应便于排查。
"""

import json
import logging
import os
import re
import tempfile
import time

import requests

logger = logging.getLogger(__name__)

WEB_HOST = "https://www.kdocs.cn"
DRIVE_HOST = "https://drive.kdocs.cn"
DEFAULT_UA = (
    "Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 "
    "(KHTML, like Gecko) Chrome/126.0.0.0 Safari/537.36"
)

RECENT_PATH = "/api/v5/links?offset=0&count=100&orderby=file_mtime&order=DESC&append=false"
GROUP_FILES_PATH = "/api/v5/groups/{gid}/files?include=acl,pic_thumbnail&with_link=true&offset=0&count=100"
DOWNLOAD_PATH = "/api/v3/groups/{gid}/files/{fid}/download?isblocks=false"

SHEET_EXTS = ["xlsx", "xls", "xlsm", "et"]
HTTP_TIMEOUT = 30

_COOKIE_FILENAME = "wps_cookie.json"


class WpsCloudError(Exception):
    """云文档操作失败。"""


class WpsAuthExpired(WpsCloudError):
    """登录凭证（wps_sid）失效或未填写。"""


def oauth_config_path():
    """Cookie 凭证文件路径：%LOCALAPPDATA%\\ExcelSendWx\\wps_cookie.json"""
    base = os.environ.get("LOCALAPPDATA") or tempfile.gettempdir()
    folder = os.path.join(base, "ExcelSendWx")
    os.makedirs(folder, exist_ok=True)
    return os.path.join(folder, _COOKIE_FILENAME)


def parse_file_input(text):
    """把用户粘贴的内容解析为文档标识。

    支持：
    - https://www.kdocs.cn/l/{短链}、365.kdocs.cn/l/{短链} 分享链接 → 返回短链 token
    - https://www.kdocs.cn/office/{id} 、/office/file/{id} 形式链接
    - 直接粘贴 ID / 短链 token
    """
    text = (text or "").strip()
    if not text:
        raise WpsCloudError("请填写在线文档 ID 或链接")

    if text.startswith("http://") or text.startswith("https://"):
        m = re.search(r"/l/([A-Za-z0-9_-]{6,})", text)
        if m:
            return m.group(1)
        m = re.search(r"/office/(?:file/)?([A-Za-z0-9_-]{6,})", text)
        if m:
            return m.group(1)
        raise WpsCloudError("无法从链接中识别文档 ID，请粘贴 kdocs.cn 文档链接或 ID")

    token = text.strip()
    # 列表选择产生的 groupid:fileid 复合标识
    if re.fullmatch(r"\d+:\d+", token):
        return token
    if not re.fullmatch(r"[A-Za-z0-9_-]{6,}", token):
        raise WpsCloudError("文档 ID 格式不正确，请检查后重试")
    return token


class WpsOAuthStore:
    """wps_sid Cookie 凭证的本地持久化（类名保留以兼容既有调用）。"""

    def __init__(self, path=None):
        self.path = path or oauth_config_path()
        self.data = {"wps_sid": "", "saved_at": 0}
        self.load()

    def load(self):
        try:
            with open(self.path, "r", encoding="utf-8") as f:
                saved = json.load(f)
            if isinstance(saved, dict):
                for key in self.data:
                    if key in saved:
                        self.data[key] = saved[key]
        except (OSError, ValueError):
            pass

    def save(self):
        self.data["saved_at"] = int(time.time())
        folder = os.path.dirname(self.path)
        os.makedirs(folder, exist_ok=True)
        fd, tmp = tempfile.mkstemp(prefix=".wps_cookie.", suffix=".tmp", dir=folder)
        try:
            with os.fdopen(fd, "w", encoding="utf-8") as f:
                json.dump(self.data, f, ensure_ascii=False, indent=2)
            os.replace(tmp, self.path)
        except Exception:
            try:
                os.remove(tmp)
            except OSError:
                pass
            raise

    @property
    def wps_sid(self):
        return (self.data.get("wps_sid") or "").strip()

    def is_configured(self):
        return bool(self.wps_sid)

    def token_fresh(self):
        return self.is_configured()

    def clear_tokens(self):
        self.data["wps_sid"] = ""
        try:
            os.remove(self.path)
        except OSError:
            pass


def _extract_ids_from_html(html, log_fn=None):
    """从网页 HTML 中提取 (group_id, file_id)。"""
    log = log_fn or (lambda msg: None)

    def _grab(*patterns):
        for pat in patterns:
            m = re.search(pat, html)
            if m:
                return m.group(1)
        return None

    gid = _grab(
        r'"group_id"\s*:\s*"?(\d+)"?',
        r'"groupid"\s*:\s*"?(\d+)"?',
        r"group_id\s*[:=]\s*['\"]?(\d+)",
    )
    fid = _grab(
        r'"file_id"\s*:\s*"?(\d+)"?',
        r'"fileid"\s*:\s*"?(\d+)"?',
        r"file_id\s*[:=]\s*['\"]?(\d+)",
    )
    if not (gid and fid):
        log("页面中未直接找到 JSON 键名，尝试宽松正则")
        gid = gid or _grab(r"group_?[iI]d[\"']?\s*[:=]\s*[\"']?(\d+)")
        fid = fid or _grab(r"file_?[iI]d[\"']?\s*[:=]\s*[\"']?(\d+)")
    return gid, fid


class WpsCloudClient:
    """带登录态调用金山文档网页版接口。"""

    def __init__(self, store: WpsOAuthStore = None, log_fn=None):
        self.store = store or WpsOAuthStore()
        self._log_fn = log_fn
        self._sid_checked = False

    def _log(self, msg):
        logger.info("%s", msg)
        if self._log_fn:
            try:
                self._log_fn(msg)
            except Exception:
                pass

    def _session(self):
        if not self.store.is_configured():
            raise WpsAuthExpired(
                "未配置登录凭证：请在“Cookie 凭证设置”中粘贴浏览器里的 wps_sid"
            )
        session = requests.Session()
        session.cookies.set("wps_sid", self.store.wps_sid, domain=".kdocs.cn")
        session.cookies.set("csrf", self.store.wps_sid, domain=".kdocs.cn")
        session.headers.update(
            {
                "User-Agent": DEFAULT_UA,
                "Origin": "https://365.kdocs.cn",
                "Referer": "https://365.kdocs.cn/",
            }
        )
        return session

    def _request(self, session, url, **kwargs):
        kwargs.setdefault("timeout", HTTP_TIMEOUT)
        resp = session.get(url, allow_redirects=True, **kwargs)
        if resp.history and any("passport" in (r.headers.get("Location") or "") for r in resp.history):
            raise WpsAuthExpired("登录凭证已失效（被重定向到登录页），请重新复制 wps_sid")
        if resp.status_code in (401, 403):
            raise WpsAuthExpired(
                f"登录凭证被拒绝（HTTP {resp.status_code}），请重新复制 wps_sid"
            )
        return resp

    # ------------------------------------------------------------------
    # 文件列表
    # ------------------------------------------------------------------

    def list_spreadsheets(self):
        """拉取最近/可用表格列表 → [{id, name, size, group_id, mtime}]。"""
        session = self._session()
        url = DRIVE_HOST + RECENT_PATH
        self._log("拉取最近文件列表：drive.kdocs.cn")
        resp = self._request(session, url)
        if resp.status_code != 200:
            raise WpsCloudError(f"文件列表接口返回 HTTP {resp.status_code}")
        try:
            payload = resp.json()
        except ValueError:
            raise WpsCloudError("文件列表接口返回非 JSON（页面可能要求重新登录）")
        files = self._parse_file_list(payload)
        self._log(f"识别到 {len(files)} 个在线表格")
        return files

    def _parse_file_list(self, payload):
        """解析 /api/v5/links 响应：文件位于 share[].file，名称在 share_name。"""
        out, seen = [], set()

        def _add(fid, name, gid="", size=0, mtime=0):
            if not fid or not name:
                return
            name = str(name)
            ext = name.rsplit(".", 1)[-1].lower() if "." in name else ""
            if ext not in SHEET_EXTS or str(fid) in seen:
                return
            seen.add(str(fid))
            out.append(
                {
                    "id": str(fid),
                    "name": name,
                    "size": int(size or 0),
                    "group_id": str(gid or ""),
                    "mtime": int(mtime or 0),
                }
            )

        shares = payload.get("share")
        if isinstance(shares, list):
            for item in shares:
                if not isinstance(item, dict):
                    continue
                f = item.get("file")
                f = f if isinstance(f, dict) else {}
                _add(
                    f.get("id"),
                    item.get("share_name") or f.get("fname"),
                    f.get("groupid") or f.get("group_id"),
                    f.get("fsize"),
                    item.get("share_ctime") or f.get("ctime"),
                )

        # 兼容 files[]/filelist[] 结构
        def _collect(node):
            if isinstance(node, dict):
                _add(
                    node.get("id") or node.get("fileid"),
                    node.get("fname") or node.get("file_name") or node.get("name"),
                    node.get("groupid") or node.get("group_id"),
                    node.get("fsize"),
                    node.get("ctime"),
                )
                for key in ("files", "filelist", "list"):
                    child = node.get(key)
                    if isinstance(child, list):
                        for sub in child:
                            _collect(sub)
            elif isinstance(node, list):
                for sub in node:
                    _collect(sub)

        _collect(payload)
        return out

    # ------------------------------------------------------------------
    # 下载
    # ------------------------------------------------------------------

    def _resolve_download_ids(self, session, token):
        """把用户输入的 token 解析为 (group_id, file_id)。"""
        page_url = f"{WEB_HOST}/l/{token}"
        self._log(f"解析文档标识：{page_url}")
        resp = self._request(session, page_url)
        html = resp.text or ""
        if "wps_sid" in (resp.request.headers.get("Cookie") or "") and (
            "账号登录" in html or "扫码登录" in html
        ):
            raise WpsAuthExpired("登录凭证已失效（页面要求登录），请重新复制 wps_sid")
        gid, fid = _extract_ids_from_html(html, self._log)
        if not (gid and fid):
            raise WpsCloudError(
                "无法从文档页面解析下载标识（网页版可能已改版）。"
                "请把日志发给开发者排查，或改用“浏览在线文档”选择。"
            )
        self._log(f"解析成功 groupid={gid} fileid={fid}")
        return gid, fid

    def download_to(self, file_token, dest_path):
        """下载在线表格为本地 xlsx。

        file_token 支持：
        - "groupid:fileid" 复合标识（浏览列表选择，直接调下载接口）
        - /l/ 短链 token 或文档 ID（先打开网页解析 groupid/fileid）
        """
        session = self._session()
        m = re.fullmatch(r"(\d+):(\d+)", file_token)
        if m:
            gid, fid = m.group(1), m.group(2)
        else:
            gid, fid = self._resolve_download_ids(session, file_token)

        api_url = DRIVE_HOST + DOWNLOAD_PATH.format(gid=gid, fid=fid)
        self._log(f"请求下载地址：{api_url}")
        resp = self._request(session, api_url)
        try:
            payload = resp.json()
        except ValueError:
            raise WpsCloudError(
                f"下载接口返回非 JSON（HTTP {resp.status_code}）：{resp.text[:200]}"
            )
        fileinfo = payload.get("fileinfo") or payload.get("data") or payload
        url = fileinfo.get("url") or fileinfo.get("download_url")
        if not url:
            raise WpsCloudError(
                f"下载接口响应中无 url 字段：{json.dumps(payload, ensure_ascii=False)[:200]}"
            )

        self._log("开始下载文件…")
        dl = session.get(url, timeout=120, stream=True)
        dl.raise_for_status()
        tmp = dest_path + ".part"
        with open(tmp, "wb") as f:
            for chunk in dl.iter_content(chunk_size=1 << 16):
                if chunk:
                    f.write(chunk)
        os.replace(tmp, dest_path)
        size = os.path.getsize(dest_path)
        self._log(f"下载完成：{size} 字节")
        return dest_path


def new_temp_xlsx_path():
    """生成 %TEMP% 下的云文档临时文件路径。"""
    folder = os.path.join(tempfile.gettempdir(), "ExcelSendWx_cloud")
    os.makedirs(folder, exist_ok=True)
    name = f"cloud_{os.getpid()}_{int(time.time() * 1000)}.xlsx"
    return os.path.join(folder, name)


def cleanup_temp_file(path):
    """读完立即清理临时文件，失败不影响主流程。"""
    if not path:
        return
    try:
        os.remove(path)
        logger.info("Cloud temp file removed: %s", path)
    except OSError:
        pass


def _calamine_workbook(file_path):
    """优先使用 Rust 实现的 calamine 引擎（大文件速度比 openpyxl 快数倍，
    且解析在原生层完成、不长时间持有 GIL，不会把界面卡成“未响应”）。
    不可用时返回 None。"""
    try:
        from python_calamine import CalamineWorkbook
        return CalamineWorkbook.from_path(file_path)
    except Exception:
        return None


def _rows_from_calamine(sheet):
    # calamine 的 None 统一转空串，与 pandas/openpyxl 分支口径一致
    return [
        [cell if cell is not None else "" for cell in row]
        for row in sheet.to_python()
    ]


def read_sheet_names(file_path, log_fn=None):
    """只读取工作簿的 sheet 名称列表（不加载数据），尽量轻量。"""
    log = log_fn or (lambda msg: None)
    wb = _calamine_workbook(file_path)
    if wb is not None:
        try:
            return list(wb.sheet_names)
        finally:
            close = getattr(wb, "close", None)
            if callable(close):
                close()
    try:
        from openpyxl import load_workbook
        wb2 = load_workbook(file_path, read_only=True, data_only=True)
        try:
            return list(wb2.sheetnames)
        finally:
            wb2.close()
    except Exception as exc:
        log(f"读取 sheet 名称失败: {exc}")
        return []


def read_one_sheet(file_path, sheet_name, log_fn=None):
    """只读取指定 sheet → 二维数组。供配置链/auto_send 使用：
    大工作簿（实测 37MB/13 sheet）全量读取要近 100 秒且吃满 GIL，
    定向读取仅需 1~4 秒。sheet 不存在时返回 None。"""
    log = log_fn or (lambda msg: None)
    wb = _calamine_workbook(file_path)
    if wb is not None:
        try:
            names = list(wb.sheet_names)
            if sheet_name not in names:
                return None
            return _rows_from_calamine(wb.get_sheet_by_name(sheet_name))
        except Exception as exc:
            log(f"calamine 读取 sheet 失败: {exc}，尝试 pandas/openpyxl…")
        finally:
            close = getattr(wb, "close", None)
            if callable(close):
                close()
    try:
        import pandas as pd
        df = pd.read_excel(file_path, sheet_name=sheet_name, header=None)
        df = df.fillna("")
        return df.values.tolist()
    except Exception as exc:
        log(f"pandas 读取 sheet 失败: {exc}，尝试 openpyxl…")
        try:
            from openpyxl import load_workbook
            wb2 = load_workbook(file_path, read_only=True, data_only=True)
            try:
                if sheet_name not in wb2.sheetnames:
                    return None
                ws = wb2[sheet_name]
                return [
                    [cell if cell is not None else "" for cell in row]
                    for row in ws.iter_rows(values_only=True)
                ]
            finally:
                wb2.close()
        except Exception as exc2:
            raise RuntimeError(f"工作表“{sheet_name}”读取失败: {exc2}") from exc


def read_all_sheets(file_path, log_fn=None):
    """读取整个工作簿所有 sheet → {sheet_name: 二维数组}。

    优先 calamine（Rust，快且不长时间占用 GIL），失败再走
    pandas/openpyxl 兜底。一次读完全部 sheet，便于下载后立刻删除临时文件。
    """
    log = log_fn or (lambda msg: None)

    wb = _calamine_workbook(file_path)
    if wb is not None:
        try:
            out = {}
            for name in wb.sheet_names:
                out[name] = _rows_from_calamine(wb.get_sheet_by_name(name))
            if out:
                return out
        except Exception as exc:
            log(f"calamine 读取失败: {exc}，尝试 pandas…")
        finally:
            close = getattr(wb, "close", None)
            if callable(close):
                close()

    def _with_openpyxl():
        from openpyxl import load_workbook
        wb = load_workbook(file_path, read_only=True, data_only=True)
        try:
            out = {}
            for ws in wb.worksheets:
                out[ws.title] = [
                    [cell if cell is not None else "" for cell in row]
                    for row in ws.iter_rows(values_only=True)
                ]
        finally:
            wb.close()
        return out

    try:
        import pandas as pd
        xls = pd.ExcelFile(file_path)
        result = {}
        try:
            for sheet_name in xls.sheet_names:
                df = pd.read_excel(xls, sheet_name=sheet_name, header=None)
                df = df.fillna("")
                result[sheet_name] = df.values.tolist()
        finally:
            xls.close()
        if result:
            return result
        log("pandas 未读到 sheet，尝试 openpyxl…")
        return _with_openpyxl()
    except Exception as exc:
        log(f"pandas 读取失败: {exc}，尝试 openpyxl…")
        return _with_openpyxl()
