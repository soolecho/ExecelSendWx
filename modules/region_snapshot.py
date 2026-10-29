# -*- coding: utf-8 -*-
"""指定区域截图：按 Excel 原始样式渲染成图片。

背景：早期「指定区域截图」是把单元格值读成二维数组后用 QPainter 重新排版
（自定义表头配色、斑马纹、自适应列宽），与原表样式差异大——合并单元格被拆开、
底色丢失、数字格式（百分比/千分位/日期/货币）被忽略，看起来像“重新画了一遍”。

本模块直接打开源工作簿读取样式（openpyxl），按原始：
  - 列宽 / 行高（隐藏行列自动跳过）
  - 合并单元格（跨行跨列不拆开，不画内部网格线）
  - 单元格填充色、字体（名称/字号/加粗/倾斜/颜色）
  - 边框样式与颜色
  - 数字格式（百分比、千分位、小数位、日期时间、货币前缀等）
  - 对齐方式 / 自动换行
绘制成 PNG。源文件不可用、渲染异常时由调用方回退到旧的重绘逻辑。

渲染纯 QImage/QPainter，不依赖 GUI，可在 worker 线程调用。
"""

import datetime as _dt
import hashlib
import logging
import os
import re
import tempfile
import time

logger = logging.getLogger(__name__)

# Excel 默认列宽 8.43 字符 ≈ 64px（Calibri 11 / 96dpi）；1 字符 ≈ 7px
_DEFAULT_COL_PX = 64
_PX_PER_CHAR = 7
_PX_PER_POINT = 96.0 / 72.0

# 输出图片上限（超宽/超高时整体等比缩小，避免生成几十兆的巨图）
_MAX_W = 3000
_MAX_H = 9000

# 主题配色：XML 里 clrScheme 的节点名（解析主题时用）
_THEME_XML_NAMES = [
    "dk1", "lt1", "dk2", "lt2",
    "accent1", "accent2", "accent3", "accent4", "accent5", "accent6",
    "hlink", "folHlink",
]
# 样式里 color/@theme 的索引顺序（0=lt1、1=dk1…）与 XML 节点顺序不同：
# Excel 默认字体就是 theme="1"（黑色 dk1），若按 XML 顺序取会变成白色。
_THEME_INDEX_ORDER = [
    "lt1", "dk1", "lt2", "dk2",
    "accent1", "accent2", "accent3", "accent4", "accent5", "accent6",
    "hlink", "folHlink",
]

# 懒渲染缓存：key -> png 路径
_PNG_CACHE = {}


# ============================== 颜色 ==============================

def _theme_palette(wb):
    """解析工作簿主题色 → {dk1: 'RRGGBB', accent1: ...}。解析失败返回 {}。"""
    raw = getattr(wb, "loaded_theme", None)
    if not raw:
        return {}
    try:
        xml = raw.decode("utf-8", "ignore") if isinstance(raw, bytes) else str(raw)
    except Exception:
        return {}
    out = {}
    for name in _THEME_XML_NAMES:
        m = re.search(
            r"<a:%s[^>]*>\s*<a:(?:srgbClr\s+val=\"([0-9A-Fa-f]{6})\""
            r"|sysClr[^>]*lastClr=\"([0-9A-Fa-f]{6})\")" % re.escape(name),
            xml,
        )
        if m:
            out[name] = (m.group(1) or m.group(2)).upper()
    return out


def _apply_tint(rgb, tint):
    """按 Excel 的 tint/shade 粗略修正亮度：tint>0 变亮，tint<0 变暗。"""
    if not tint:
        return rgb
    r = int(rgb[0:2], 16)
    g = int(rgb[2:4], 16)
    b = int(rgb[4:6], 16)
    if tint > 0:
        r = int(r + (255 - r) * tint)
        g = int(g + (255 - g) * tint)
        b = int(b + (255 - b) * tint)
    else:
        k = 1.0 + tint
        r, g, b = int(r * k), int(g * k), int(b * k)
    return "%02X%02X%02X" % (
        max(0, min(255, r)), max(0, min(255, g)), max(0, min(255, b))
    )


def _resolve_rgb(color, theme, default=None):
    """openpyxl 颜色对象 → 'RRGGBB'；无法解析返回 default。"""
    if color is None:
        return default
    if isinstance(color, str):
        text = color.strip()
        if len(text) == 8:
            text = text[2:]
        return text.upper() if re.fullmatch(r"[0-9A-Fa-f]{6}", text or "") else default
    ctype = getattr(color, "type", None)
    try:
        if ctype == "rgb":
            rgb = getattr(color, "rgb", None) or ""
            return _resolve_rgb(rgb, theme, default)
        if ctype == "indexed":
            from openpyxl.styles.colors import COLOR_INDEX
            idx = getattr(color, "indexed", None)
            if idx is not None and 0 <= int(idx) < len(COLOR_INDEX):
                return _resolve_rgb(COLOR_INDEX[int(idx)], theme, default)
            return default
        if ctype == "theme":
            idx = getattr(color, "theme", None)
            name = (_THEME_INDEX_ORDER[int(idx)]
                    if idx is not None and int(idx) < len(_THEME_INDEX_ORDER)
                    else "")
            rgb = theme.get(name)
            if not rgb:
                return default
            return _apply_tint(rgb, getattr(color, "tint", 0.0))
        if ctype == "auto":
            return default
    except Exception:
        return default
    return default


def _qcolor(rgb, fallback=None):
    from PyQt6.QtGui import QColor
    if not rgb:
        return QColor(fallback) if fallback else None
    return QColor("#" + rgb)


# ============================== 数字格式 ==============================

def _split_sections(fmt):
    """按分号拆格式段（引号内的分号不算）。"""
    sections, cur, in_quote = [], "", False
    for ch in fmt:
        if ch == '"':
            in_quote = not in_quote
            cur += ch
            continue
        if ch == ";" and not in_quote:
            sections.append(cur)
            cur = ""
            continue
        cur += ch
    sections.append(cur)
    return sections


def _strip_noise(section):
    """去掉 [Red]/[$¥-804]/_ 占位/\\转义 等修饰，保留可见文本。"""
    out = []
    i, n = 0, len(section)
    while i < n:
        ch = section[i]
        if ch == '"':
            j = section.find('"', i + 1)
            if j < 0:
                out.append(section[i + 1:])
                break
            out.append(section[i + 1:j])
            i = j + 1
            continue
        if ch == "\\" and i + 1 < n:
            out.append(section[i + 1])
            i += 2
            continue
        if ch == "_" and i + 1 < n:
            # Excel 用 _x 预留字符宽度；近似为一个半角空格
            out.append(" ")
            i += 2
            continue
        if ch == "*" and i + 1 < n:
            i += 2
            continue
        if ch == "[":
            j = section.find("]", i)
            if j < 0:
                break
            inner = section[i + 1:j]
            m = re.match(r"\$([^\-]*)", inner)
            if m and m.group(1):
                out.append(m.group(1))  # [$¥-804] → ¥
            i = j + 1
            continue
        out.append(ch)
        i += 1
    return "".join(out)


def _general_number(value):
    """General 格式：整数直出，浮点去掉多余尾零（最多 10 位有效数字）。"""
    if isinstance(value, int):
        return str(value)
    f = float(value)
    if f.is_integer() and abs(f) < 1e15:
        return str(int(f))
    text = "%.10g" % f
    return text


def _fmt_numeric(value, fmt):
    text_fmt = (fmt or "").strip()
    if not text_fmt or text_fmt.lower() == "general":
        return _general_number(value)

    sections = _split_sections(text_fmt)
    use_neg_section = False
    if value < 0 and len(sections) >= 2:
        section = sections[1]
        use_neg_section = True
    elif value == 0 and len(sections) >= 3:
        section = sections[2]
    else:
        section = sections[0]

    if "@" in section and section.strip() == "@":
        return _general_number(value)

    cleaned = _strip_noise(section)
    m = re.search(r"[#?0][#?0,]*(?:\.[#?0]+)?(?:[Ee][+\-]?[#0]+)?", cleaned)
    if not m:
        # 该格式段没有数字占位符（如 0;-0;"-" 的零值段）→ 直接用字面文本
        return cleaned if cleaned.strip() else _general_number(value)

    pattern = m.group(0)
    prefix, suffix = cleaned[:m.start()], cleaned[m.end():]
    av = abs(float(value))
    is_percent = "%" in (prefix + suffix)
    if is_percent:
        av *= 100.0

    dec_part = pattern.split(".", 1)[1] if "." in pattern else ""
    dec_part = re.split(r"[Ee]", dec_part)[0]
    decimals = sum(1 for ch in dec_part if ch in "0#?")
    optional_dec = "#" in dec_part or "?" in dec_part
    thousands = "," in pattern
    scientific = bool(re.search(r"[Ee][+\-]", pattern))

    try:
        if scientific:
            body = ("%." + str(decimals) + "E") % av
        elif thousands:
            body = ("{:,.%df}" % decimals).format(av)
        else:
            body = ("%." + str(decimals) + "f") % av
    except (ValueError, TypeError):
        return _general_number(value)

    if optional_dec and "." in body:
        body = body.rstrip("0").rstrip(".")
    if thousands and len(sections) >= 1 and value < 0:
        pass

    sign = ""
    if value < 0 and not use_neg_section and "(" not in prefix:
        sign = "-"
    return "%s%s%s%s" % (prefix, sign, body, suffix)


_WEEKDAYS = ["Monday", "Tuesday", "Wednesday", "Thursday", "Friday",
             "Saturday", "Sunday"]
_MONTHS = ["January", "February", "March", "April", "May", "June", "July",
           "August", "September", "October", "November", "December"]


def _date_tokenize(fmt):
    """日期格式 → [('lit', 文本) / ('tok', 'yyyy') 等]。"""
    tokens = []
    i, n = 0, len(fmt)
    while i < n:
        ch = fmt[i]
        if ch == '"':
            j = fmt.find('"', i + 1)
            if j < 0:
                tokens.append(("lit", fmt[i + 1:]))
                break
            tokens.append(("lit", fmt[i + 1:j]))
            i = j + 1
            continue
        if ch == "\\" and i + 1 < n:
            tokens.append(("lit", fmt[i + 1]))
            i += 2
            continue
        if ch == "[":
            j = fmt.find("]", i)
            i = n if j < 0 else j + 1
            continue
        if ch in "aA" and fmt[i:i + 5].upper() == "AM/PM":
            tokens.append(("ampm", ""))
            i += 5
            continue
        if ch in "yYmdhHsS":
            j = i
            while j < n and fmt[j] in "yYmdhHsS":
                j += 1
            tokens.append(("tok", fmt[i:j]))
            i = j
            continue
        if ch in "0#":
            j = i
            while j < n and fmt[j] in "0#":
                j += 1
            tokens.append(("lit", fmt[i:j]))
            i = j
            continue
        tokens.append(("lit", ch))
        i += 1
    return tokens


def _fmt_datetime(value, fmt):
    text_fmt = (fmt or "").strip()
    is_time_only = isinstance(value, _dt.time)
    if not text_fmt or text_fmt.lower() == "general":
        if is_time_only:
            text_fmt = "h:mm:ss"
        elif isinstance(value, _dt.datetime) and value.time() != _dt.time(0, 0):
            text_fmt = "yyyy-mm-dd hh:mm:ss"
        else:
            text_fmt = "yyyy-mm-dd"

    # time 对象没有年月日，补成同一天的 datetime 以便统一处理
    if is_time_only:
        dt = _dt.datetime.combine(_dt.date(1899, 12, 30), value)
    else:
        dt = value

    tokens = _date_tokenize(text_fmt)
    has_ampm = any(k == "ampm" for k, _v in tokens)
    out = []
    for idx, (kind, tok) in enumerate(tokens):
        if kind == "lit":
            out.append(tok)
            continue
        if kind == "ampm":
            out.append("AM" if dt.hour < 12 else "PM")
            continue
        low = tok.lower()
        prev_tok = ""
        for k, v in reversed(tokens[:idx]):
            if k == "tok":
                prev_tok = v.lower()
                break
        next_tok = ""
        for k, v in tokens[idx + 1:]:
            if k == "tok":
                next_tok = v.lower()
                break
        # m / mm 需要区分「月」还是「分钟」：紧跟 h/H 或紧接 s/S → 分钟
        is_minute = low[0] == "m" and (
            prev_tok.startswith("h") or next_tok.startswith("s")
            or (is_time_only and not prev_tok and not next_tok)
        )
        if low[0] == "y":
            if len(low) >= 4:
                out.append("%04d" % dt.year)
            elif len(low) == 2:
                out.append("%02d" % (dt.year % 100))
            else:
                out.append(str(dt.year % 100))
        elif low[0] == "m" and is_minute:
            out.append(("%02d" if len(low) >= 2 else "%d") % dt.minute)
        elif low[0] == "m":
            if len(low) == 1:
                out.append(str(dt.month))
            elif len(low) == 2:
                out.append("%02d" % dt.month)
            elif len(low) == 3:
                out.append(_MONTHS[dt.month - 1][:3])
            else:
                out.append(_MONTHS[dt.month - 1])
        elif low[0] == "d":
            if len(low) == 1:
                out.append(str(dt.day))
            elif len(low) == 2:
                out.append("%02d" % dt.day)
            elif len(low) == 3:
                out.append(_WEEKDAYS[dt.weekday()][:3])
            else:
                out.append(_WEEKDAYS[dt.weekday()])
        elif low[0] == "h":
            hour = dt.hour
            if has_ampm:
                hour = hour % 12 or 12
            out.append(("%02d" if len(low) >= 2 else "%d") % hour)
        elif low[0] == "s":
            out.append(("%02d" if len(low) >= 2 else "%d") % dt.second)
        else:
            out.append(tok)
    text = "".join(out).strip()
    return text


def _fmt_timedelta(value, fmt):
    total = int(value.total_seconds())
    hours, rem = divmod(total, 3600)
    minutes, seconds = divmod(rem, 60)
    text_fmt = (fmt or "").strip().lower()
    if "mm" in text_fmt and "hh" not in text_fmt:
        return "%d:%02d" % (total // 60, total % 60)
    if "ss" not in text_fmt:
        return "%d:%02d" % (hours, minutes)
    return "%d:%02d:%02d" % (hours, minutes, seconds)


def format_excel_value(value, number_format=""):
    """按单元格数字格式把原始值渲染成 Excel 里显示的文本。"""
    if value is None:
        return ""
    if isinstance(value, str):
        return value
    if isinstance(value, bool):
        return "TRUE" if value else "FALSE"
    try:
        if isinstance(value, _dt.timedelta):
            return _fmt_timedelta(value, number_format)
        if isinstance(value, _dt.datetime):
            return _fmt_datetime(value, number_format)
        if isinstance(value, _dt.date):
            return _fmt_datetime(value, number_format)
        if isinstance(value, _dt.time):
            return _fmt_datetime(value, number_format)
        if isinstance(value, (int, float)):
            return _fmt_numeric(value, number_format)
    except Exception:
        logger.debug("单元格格式化失败: %r / %r", value, number_format)
    return str(value)


# ============================== 几何 ==============================

def _col_px(width_chars):
    if not width_chars:
        return _DEFAULT_COL_PX
    try:
        return max(8, int(round(float(width_chars) * _PX_PER_CHAR + 5)))
    except (TypeError, ValueError):
        return _DEFAULT_COL_PX


def _row_px(height_pt):
    if not height_pt:
        return int(round(15 * _PX_PER_POINT))
    try:
        return max(6, int(round(float(height_pt) * _PX_PER_POINT)))
    except (TypeError, ValueError):
        return int(round(15 * _PX_PER_POINT))


_BORDER_STYLES = {
    "thin": 1.0,
    "medium": 2.0,
    "thick": 3.0,
    "double": 2.0,
    "hair": 0.6,
    "dotted": 1.0,
    "dashed": 1.0,
    "dashDot": 1.0,
    "dashDotDot": 1.0,
    "mediumDashed": 2.0,
    "mediumDashDot": 2.0,
    "mediumDashDotDot": 2.0,
    "slantDashDot": 2.0,
}

_DASH_STYLES = {"dotted": "DotLine", "hair": "DotLine", "dashed": "DashLine",
                "dashDot": "DashLine", "dashDotDot": "DashLine",
                "mediumDashed": "DashLine", "mediumDashDot": "DashLine",
                "mediumDashDotDot": "DashLine", "slantDashDot": "DashLine"}


def _side_pen(side, theme):
    """openpyxl 边框 side → (宽度, 虚线样式, 颜色)。"""
    style = getattr(side, "style", None) if side is not None else None
    if not style:
        return None
    width = _BORDER_STYLES.get(style, 1.0)
    dash = _DASH_STYLES.get(style)
    rgb = _resolve_rgb(getattr(side, "color", None), theme, default="000000")
    return width, dash, rgb


def _fill_rgb(cell, theme):
    fill = getattr(cell, "fill", None)
    if fill is None:
        return None
    pattern = getattr(fill, "patternType", None) or getattr(fill, "fill_type", None)
    if not pattern or pattern in ("none", "gray125"):
        # GradientFill 没有 patternType
        stops = getattr(fill, "stop", None)
        if stops:
            try:
                return _resolve_rgb(stops[0].color, theme)
            except Exception:
                return None
        return None
    return _resolve_rgb(getattr(fill, "fgColor", None), theme)


# ============================== 条件格式 ==============================
# Excel 里「颜色凸显」很多是条件格式（色阶/单元格规则/前 N 项）而不是单元格
# 直接填充，openpyxl 的 cell.fill 读不到，需要单独解析 ws.conditional_formatting。

def _to_number(value):
    if isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    if isinstance(value, str):
        try:
            return float(value.strip())
        except (TypeError, ValueError):
            return None
    return None


def _formula_value(ws, formula):
    """条件格式公式里的比较值：数字 / "文本" / $A$1 引用。"""
    if formula is None:
        return None
    text = str(formula).strip()
    if len(text) >= 2 and text[0] == '"' and text[-1] == '"':
        return text[1:-1]
    num = _to_number(text)
    if num is not None:
        return num
    m = re.fullmatch(r"\$?([A-Za-z]{1,3})\$?(\d+)", text)
    if m:
        try:
            return ws[f"{m.group(1).upper()}{m.group(2)}"].value
        except Exception:
            return None
    return None


def _dxf_fill_rgb(fill, theme):
    """差分样式（条件格式）的填充色：Excel 常写在 bgColor 上。"""
    if fill is None:
        return None
    rgb = _resolve_rgb(getattr(fill, "bgColor", None), theme)
    if not rgb:
        rgb = _resolve_rgb(getattr(fill, "fgColor", None), theme)
    return rgb


def _dxf_style(dxf, theme):
    style = {}
    if dxf is None:
        return style
    try:
        rgb = _dxf_fill_rgb(getattr(dxf, "fill", None), theme)
        if rgb:
            style["fill"] = rgb
    except Exception:
        pass
    font = getattr(dxf, "font", None)
    if font is not None:
        rgb = _resolve_rgb(getattr(font, "color", None), theme)
        if rgb:
            style["font"] = rgb
        if getattr(font, "bold", None):
            style["bold"] = True
    return style


def _merge_style(result, cell, style):
    """多条规则命中同一单元格时，优先级高的（先出现的）不被覆盖。"""
    cur = result.setdefault(cell, {})
    for key, value in style.items():
        if value is not None and key not in cur:
            cur[key] = value


def _cell_value(ws, row, col):
    cells = getattr(ws, "_cells", None)
    if cells is not None:
        cell = cells.get((row, col))
        return getattr(cell, "value", None) if cell is not None else None
    return getattr(ws.cell(row=row, column=col), "value", None)


def _range_values(ws, rng):
    values = []
    for row in range(rng.min_row, rng.max_row + 1):
        for col in range(rng.min_col, rng.max_col + 1):
            values.append(_cell_value(ws, row, col))
    return values


def _interp_color(stops, value):
    """按色阶停靠点线性插值 → 'RRGGBB'。"""
    if not stops:
        return None
    if value <= stops[0][0]:
        return stops[0][1]
    if value >= stops[-1][0]:
        return stops[-1][1]
    for (v0, c0), (v1, c1) in zip(stops, stops[1:]):
        if v0 <= value <= v1:
            if v1 == v0:
                return c0
            ratio = (value - v0) / float(v1 - v0)
            rgb = []
            for i in (0, 2, 4):
                a = int(c0[i:i + 2], 16)
                b = int(c1[i:i + 2], 16)
                rgb.append("%02X" % max(0, min(255, int(round(a + (b - a) * ratio)))))
            return "".join(rgb)
    return stops[-1][1]


def _color_scale_cells(ws, ranges, rule, theme):
    """色阶规则 → {(row, col): 'RRGGBB'}（按区域内最大/最小值插值）。"""
    scale = getattr(rule, "colorScale", None)
    if scale is None:
        return {}
    colors = [_resolve_rgb(c, theme) for c in (getattr(scale, "color", None) or [])]
    cfvos = list(getattr(scale, "cfvo", None) or [])
    if len(colors) < 2 or len(cfvos) < 2 or len(cfvos) != len(colors):
        return {}

    pairs, numbers = [], []
    for rng in ranges:
        for row in range(rng.min_row, rng.max_row + 1):
            for col in range(rng.min_col, rng.max_col + 1):
                num = _to_number(_cell_value(ws, row, col))
                if num is not None:
                    pairs.append(((row, col), num))
                    numbers.append(num)
    if not numbers:
        return {}

    vmin, vmax = min(numbers), max(numbers)
    sorted_nums = sorted(numbers)
    stops = []
    for idx, cfvo in enumerate(cfvos):
        ctype = (getattr(cfvo, "type", None) or "min").lower()
        raw = getattr(cfvo, "val", None)
        if ctype == "min":
            pos = vmin
        elif ctype == "max":
            pos = vmax
        elif ctype == "num":
            try:
                pos = float(raw)
            except (TypeError, ValueError):
                pos = vmin if idx == 0 else vmax
        elif ctype == "percent":
            try:
                pos = vmin + (vmax - vmin) * float(raw) / 100.0
            except (TypeError, ValueError):
                pos = vmin if idx == 0 else vmax
        elif ctype == "percentile":
            try:
                k = int(round(float(raw) / 100.0 * (len(sorted_nums) - 1)))
            except (TypeError, ValueError):
                k = 0
            pos = sorted_nums[max(0, min(len(sorted_nums) - 1, k))]
        else:
            return {}
        stops.append((pos, colors[idx]))
    stops.sort(key=lambda item: item[0])

    return {cell: _interp_color(stops, num) for cell, num in pairs}


def _rule_matches(rule, rtype, value, stats):
    """单元格是否命中条件格式规则（只覆盖常见类型，其余忽略）。"""
    num = _to_number(value)
    text = "" if value is None else str(value)

    if rtype == "cellis":
        op = (getattr(rule, "operator", None) or "").lower()
        numbers = [_to_number(_formula_value(None, f))
                   for f in (getattr(rule, "formula", None) or [])]
        if num is None or not numbers or numbers[0] is None:
            return False
        first = numbers[0]
        second = numbers[1] if len(numbers) > 1 else None
        if op == "greaterthan":
            return num > first
        if op == "lessthan":
            return num < first
        if op == "greaterthanorequal":
            return num >= first
        if op == "lessthanorequal":
            return num <= first
        if op == "equal":
            return num == first
        if op == "notequal":
            return num != first
        if op == "between" and second is not None:
            return first <= num <= second
        if op == "notbetween" and second is not None:
            return not (first <= num <= second)
        return False

    if rtype in ("containstext", "notcontainstext", "beginswith", "endswith"):
        target = str(_formula_value(None, (getattr(rule, "formula", None) or [""])[0]) or "")
        if not target:
            return False
        if rtype == "containstext":
            return target in text
        if rtype == "notcontainstext":
            return target not in text
        if rtype == "beginswith":
            return text.startswith(target)
        return text.endswith(target)

    numbers = [n for n in (_to_number(v) for v in (stats or [])) if n is not None]
    if rtype == "aboveaverage":
        if num is None or not numbers:
            return False
        return num > sum(numbers) / float(len(numbers))
    if rtype == "top10":
        if num is None or not numbers:
            return False
        rank = getattr(rule, "rank", 10) or 10
        try:
            rank = float(rank)
        except (TypeError, ValueError):
            rank = 10.0
        bottom = bool(getattr(rule, "bottom", False))
        if getattr(rule, "percent", False):
            rank = max(1.0, len(numbers) * rank / 100.0)
        k = max(1, min(len(numbers), int(rank)))
        ranked = sorted(numbers)
        return num <= ranked[k - 1] if bottom else num >= ranked[-k]
    return False


def _conditional_style_map(ws, theme):
    """解析工作表条件格式 → {(row, col): {fill/font/bold}}。"""
    result = {}
    try:
        conds = list(ws.conditional_formatting)
    except Exception:
        return result
    if not conds:
        return result

    entries = []
    for cf in conds:
        for rule in cf.rules:
            prio = getattr(rule, "priority", None)
            entries.append((prio if prio is not None else 1, cf, rule))
    entries.sort(key=lambda item: item[0])

    for _prio, cf, rule in entries:
        try:
            rtype = (rule.type or "").lower()
            ranges = list(cf.sqref.ranges)
            if rtype == "colorscale":
                for cell, rgb in _color_scale_cells(ws, ranges, rule, theme).items():
                    if rgb:
                        _merge_style(result, cell, {"fill": rgb})
                continue
            style = _dxf_style(getattr(rule, "dxf", None), theme)
            if not style:
                continue
            for rng in ranges:
                stats = _range_values(ws, rng)
                for row in range(rng.min_row, rng.max_row + 1):
                    for col in range(rng.min_col, rng.max_col + 1):
                        value = _cell_value(ws, row, col)
                        if value is None:
                            continue
                        if _rule_matches(rule, rtype, value, stats):
                            _merge_style(result, (row, col), style)
        except Exception:
            logger.debug("条件格式解析失败", exc_info=True)
            continue
    return result


# ============================== 渲染 ==============================

def _used_bounds(ws):
    """单元格中真正有值的范围（0-based 半开区间）；无值返回 None。"""
    min_r = min_c = None
    max_r = max_c = -1
    cells = getattr(ws, "_cells", None) or {}
    for (row, col), cell in cells.items():
        value = getattr(cell, "value", None)
        if value is None or (isinstance(value, str) and not value.strip()):
            continue
        r, c = row - 1, col - 1
        min_r = r if min_r is None or r < min_r else min_r
        min_c = c if min_c is None or c < min_c else min_c
        max_r = max(max_r, r)
        max_c = max(max_c, c)
    if min_r is None:
        return None
    return min_r, max_r + 1, min_c, max_c + 1


def render_region_image(file_path, sheet_name="", spec="", out_png="",
                        log_fn=None):
    """把工作簿 sheet 的指定区域按原始样式渲染成 PNG。

    spec 为 "A1:F20"（空/非法 = 该 sheet 已用区域）。
    返回生成路径；源文件不可用/区域无内容/渲染异常返回 None。
    """
    log = log_fn or (lambda _m: None)
    if not file_path or not os.path.isfile(file_path):
        return None
    if file_path.lower().endswith(".xls"):
        log("区域截图：.xls 旧格式无法读取样式，改用重绘表格")
        return None

    from .monitor_engine import parse_region

    try:
        from openpyxl import load_workbook
        from openpyxl.utils import get_column_letter
    except Exception as exc:
        log(f"区域截图渲染不可用（openpyxl 导入失败）: {exc}")
        return None

    t0 = time.time()
    try:
        wb = load_workbook(file_path, data_only=True)
    except Exception as exc:
        log(f"区域截图打开工作簿失败: {exc}")
        return None

    try:
        want_sheet = (sheet_name or "").strip()
        if want_sheet and want_sheet in wb.sheetnames:
            ws = wb[want_sheet]
        else:
            if want_sheet:
                log(f"⚠ 工作表「{want_sheet}」不存在，区域截图改用「{wb.sheetnames[0]}」")
            ws = wb.worksheets[0]

        region = parse_region(spec)
        if region:
            r0, r1, c0, c1 = region
        else:
            bounds = _used_bounds(ws)
            if not bounds:
                log("区域截图：工作表没有有效内容")
                return None
            r0, r1, c0, c1 = bounds

        r0 = max(0, r0)
        c0 = max(0, c0)
        r1 = min(r1, max(ws.max_row or 1, r1))
        c1 = min(c1, max(ws.max_column or 1, c1))
        if r1 <= r0 or c1 <= c0:
            return None

        theme = _theme_palette(wb)
        return _draw(ws, theme, r0, r1, c0, c1, out_png, log,
                     get_column_letter, t0)
    finally:
        try:
            wb.close()
        except Exception:
            pass


def _draw(ws, theme, r0, r1, c0, c1, out_png, log, get_column_letter, t0):
    from PyQt6.QtCore import QPointF, QRectF, Qt
    from PyQt6.QtGui import (QColor, QFont, QImage, QPainter, QPen)

    # ---------- 列宽 / 行高（隐藏行列跳过，与 Excel 截图一致） ----------
    col_dims = {}
    for dim in ws.column_dimensions.values():
        lo = getattr(dim, "min", None) or 0
        hi = getattr(dim, "max", None) or lo
        for idx in range(lo, hi + 1):
            col_dims[idx] = dim

    cols = []  # [(1-based 列号, 像素宽)]
    for c in range(c0 + 1, c1 + 1):
        dim = col_dims.get(c)
        if dim is not None and getattr(dim, "hidden", False):
            continue
        width = getattr(dim, "width", None) if dim is not None else None
        cols.append((c, _col_px(width)))
    if not cols:
        return None

    rows = []  # [(1-based 行号, 像素高)]
    for r in range(r0 + 1, r1 + 1):
        dim = ws.row_dimensions.get(r) if hasattr(ws.row_dimensions, "get") else None
        if dim is not None and getattr(dim, "hidden", False):
            continue
        height = getattr(dim, "height", None) if dim is not None else None
        rows.append((r, _row_px(height)))
    if not rows:
        return None

    col_x, total_w = {}, 0
    for c, w in cols:
        col_x[c] = total_w
        total_w += w
    row_y, total_h = {}, 0
    for r, h in rows:
        row_y[r] = total_h
        total_h += h
    col_w = dict(cols)
    row_h = dict(rows)

    # ---------- 合并单元格 ----------
    owner, span = {}, {}
    for rng in ws.merged_cells.ranges:
        mr0, mr1 = rng.min_row, rng.max_row
        mc0, mc1 = rng.min_col, rng.max_col
        if mr1 < r0 + 1 or mr0 > r1 or mc1 < c0 + 1 or mc0 > c1:
            continue
        for rr in range(mr0, mr1 + 1):
            for cc in range(mc0, mc1 + 1):
                owner[(rr, cc)] = (mr0, mc0)
        span[(mr0, mc0)] = (mr1, mc1)

    def _cell_rect(r, c):
        """普通单元格矩形（不在区域内返回 None）。"""
        if r not in row_y or c not in col_x:
            return None
        return (col_x[c], row_y[r], col_w[c], row_h[r])

    def _anchor_rect(r, c):
        """合并区域矩形，裁剪到当前截图区域内。"""
        max_r, max_c = span.get((r, c), (r, c))
        x = col_x.get(c)
        y = row_y.get(r)
        if x is None or y is None:
            return None
        # 结束边：区域内最后一个可见列/行
        x_end = x
        for cc, w in cols:
            if c <= cc <= max_c:
                x_end = col_x[cc] + w
        y_end = y
        for rr, h in rows:
            if r <= rr <= max_r:
                y_end = row_y[rr] + h
        return (x, y, max(0, x_end - x), max(0, y_end - y))

    scale = 1.0
    if total_w > _MAX_W:
        scale = _MAX_W / float(total_w)
    if total_h * scale > _MAX_H:
        scale = _MAX_H / float(total_h)
    scale = max(0.15, min(1.0, scale))

    img_w = max(1, int(round(total_w * scale)))
    img_h = max(1, int(round(total_h * scale)))
    img = QImage(img_w, img_h, QImage.Format.Format_ARGB32)
    img.fill(QColor("#FFFFFF"))

    p = QPainter(img)
    try:
        p.setRenderHint(QPainter.RenderHint.TextAntialiasing, True)
        p.scale(scale, scale)

        # ---------- 1) 底色（单元格直接填充 + 条件格式底色） ----------
        cond_styles = _conditional_style_map(ws, theme)
        for r, _h in rows:
            for c, _w in cols:
                anch = owner.get((r, c))
                if anch and anch != (r, c):
                    continue
                rect = _anchor_rect(r, c)
                if not rect:
                    continue
                cell = ws.cell(row=r, column=c)
                cond = cond_styles.get((r, c)) or {}
                rgb = cond.get("fill") or _fill_rgb(cell, theme)
                if rgb:
                    color = _qcolor(rgb)
                    if color is not None:
                        p.fillRect(QRectF(*rect), color)

        # ---------- 2) 网格线（工作表显示网格线时才有；合并区域内部不画） ----------
        show_grid = True
        view = getattr(ws, "sheet_view", None)
        if view is not None and getattr(view, "showGridLines", None) is not None:
            show_grid = bool(view.showGridLines)
        if show_grid:
            p.setPen(QPen(QColor("#D9D9D9"), 0.8))

            def _anchor_of(r, c):
                return owner.get((r, c)) or (r, c)

            for r, _h in rows:
                y0 = row_y[r]
                for c, w in cols:
                    # 右侧边线：同一合并区域内部的两列之间不画
                    if c + 1 in col_x:
                        if _anchor_of(r, c) != _anchor_of(r, c + 1):
                            x = col_x[c] + w
                            p.drawLine(QPointF(x, y0), QPointF(x, y0 + row_h[r]))
                    # 下侧边线：同一合并区域内部的两行之间不画
                    if r + 1 in row_y:
                        if _anchor_of(r, c) != _anchor_of(r + 1, c):
                            y = y0 + row_h[r]
                            p.drawLine(QPointF(col_x[c], y),
                                       QPointF(col_x[c] + w, y))
            # 截图区域最外圈补一条（裁剪边界处 Excel 仍有网格线）
            p.drawLine(QPointF(total_w, 0), QPointF(total_w, total_h))
            p.drawLine(QPointF(0, total_h), QPointF(total_w, total_h))

        # ---------- 3) 显式边框（合并区域只画外框） ----------
        for r, _h in rows:
            for c, _w in cols:
                anch = owner.get((r, c))
                if anch and anch != (r, c):
                    continue
                rect = _anchor_rect(r, c)
                if not rect:
                    continue
                x, y, w, h = rect
                rng = span.get((r, c), (r, c))
                max_r, max_c = rng
                cell = ws.cell(row=r, column=c)
                border = getattr(cell, "border", None)
                sides = {}
                if border is not None:
                    sides["top"] = border.top
                    sides["left"] = border.left
                    right_cell = ws.cell(row=r, column=max_c)
                    bottom_cell = ws.cell(row=max_r, column=c)
                    sides["right"] = getattr(
                        getattr(right_cell, "border", None), "right", border.right
                    )
                    sides["bottom"] = getattr(
                        getattr(bottom_cell, "border", None), "bottom", border.bottom
                    )
                lines = {
                    "top": ((x, y), (x + w, y)),
                    "bottom": ((x, y + h), (x + w, y + h)),
                    "left": ((x, y), (x, y + h)),
                    "right": ((x + w, y), (x + w, y + h)),
                }
                for side_name, side in sides.items():
                    info = _side_pen(side, theme)
                    if not info:
                        continue
                    width, dash, rgb = info
                    pen = QPen(_qcolor(rgb, "#000000"), width)
                    if dash:
                        pen.setStyle(getattr(Qt.PenStyle, dash))
                    p.setPen(pen)
                    a, b = lines[side_name]
                    p.drawLine(QPointF(*a), QPointF(*b))

        # ---------- 4) 文字 ----------
        for r, _h in rows:
            for c, _w in cols:
                anch = owner.get((r, c))
                if anch and anch != (r, c):
                    continue
                rect = _anchor_rect(r, c)
                if not rect:
                    continue
                cell = ws.cell(row=r, column=c)
                text = format_excel_value(
                    getattr(cell, "value", None),
                    str(getattr(cell, "number_format", "") or ""),
                )
                if not text:
                    continue
                font_obj = getattr(cell, "font", None)
                size_pt = getattr(font_obj, "size", None) or 11
                font_name = getattr(font_obj, "name", None) or "Microsoft YaHei"
                font = QFont()
                # 中英文字体名可能对不上（如 Excel 写“微软雅黑”，Qt 只认
                # Microsoft YaHei），用字体族回退链保证字形正常
                font.setFamilies([font_name, "Microsoft YaHei", "SimSun",
                                  "Segoe UI", "Calibri"])
                font.setPixelSize(max(6, int(round(float(size_pt) * _PX_PER_POINT))))
                cond = cond_styles.get((r, c)) or {}
                font.setBold(bool(getattr(font_obj, "bold", False))
                             or bool(cond.get("bold")))
                font.setItalic(bool(getattr(font_obj, "italic", False)))
                rgb = (cond.get("font")
                       or _resolve_rgb(getattr(font_obj, "color", None), theme,
                                       default="000000"))
                p.setFont(font)
                p.setPen(QPen(_qcolor(rgb, "#000000")))

                align = getattr(cell, "alignment", None)
                horizontal = (getattr(align, "horizontal", None) or "general")
                vertical = (getattr(align, "vertical", None) or "bottom")
                wrap = bool(getattr(align, "wrap_text", False))

                flags = Qt.AlignmentFlag(0)
                if horizontal in ("center", "centerContinuous"):
                    flags |= Qt.AlignmentFlag.AlignHCenter
                elif horizontal in ("right", "distributed"):
                    flags |= Qt.AlignmentFlag.AlignRight
                elif horizontal == "justify":
                    flags |= Qt.AlignmentFlag.AlignLeft
                else:
                    # general：数字靠右，文本靠左
                    if isinstance(getattr(cell, "value", None), (int, float)) and not isinstance(
                        getattr(cell, "value", None), bool
                    ):
                        flags |= Qt.AlignmentFlag.AlignRight
                    else:
                        flags |= Qt.AlignmentFlag.AlignLeft

                if vertical == "top":
                    flags |= Qt.AlignmentFlag.AlignTop
                elif vertical in ("center", "justify", "distributed"):
                    flags |= Qt.AlignmentFlag.AlignVCenter
                else:
                    flags |= Qt.AlignmentFlag.AlignBottom
                if wrap:
                    flags |= Qt.TextFlag.TextWordWrap

                pad_x, pad_y = 3, 1
                x, y, w, h = rect

                # 不换行的左对齐文字，右侧相邻单元格为空时允许溢出（与 Excel
                # 一致），避免长文本被硬裁掉一半
                if not wrap and (flags & (Qt.AlignmentFlag.AlignHCenter
                                          | Qt.AlignmentFlag.AlignRight)) == 0:
                    from PyQt6.QtGui import QFontMetrics
                    need = QFontMetrics(font).horizontalAdvance(text) + pad_x * 2
                    extra, cursor = 0, c
                    while need > w + extra and cursor < c1:
                        nxt = cursor + 1
                        if nxt not in col_x or owner.get((r, nxt)) is not None:
                            break
                        nxt_val = getattr(ws.cell(row=r, column=nxt), "value", None)
                        if nxt_val not in (None, "") and str(nxt_val).strip():
                            break
                        extra += col_w[nxt]
                        cursor = nxt
                    w += extra

                text_rect = QRectF(x + pad_x, y + pad_y,
                                   max(1, w - pad_x * 2),
                                   max(1, h - pad_y * 2))
                p.drawText(text_rect, int(flags), text)
    finally:
        p.end()

    out_png = out_png or os.path.join(
        tempfile.gettempdir(),
        "region_snap_%s.png" % hashlib.md5(
            ("%s|%s" % (r0, c0)).encode("utf-8")).hexdigest()[:10],
    )
    try:
        os.makedirs(os.path.dirname(out_png) or ".", exist_ok=True)
    except OSError:
        pass
    if not img.save(out_png, "PNG"):
        log("区域截图保存失败")
        return None
    log(f"区域截图已按原表样式渲染（{len(rows)} 行 × {len(cols)} 列，"
        f"{time.time() - t0:.1f}秒）")
    return out_png


# ============================== 懒渲染（发送链路） ==============================

def _spec_cache_key(spec):
    """渲染规格 → 缓存 key（源文件/云文档 + sheet + 区域 + mtime）。"""
    source_path = str(spec.get("source_path") or "").strip()
    cloud_id = str(spec.get("cloud_file_id") or "").strip()
    sheet = str(spec.get("sheet") or "").strip()
    region = str(spec.get("range") or "").strip()
    mtime = spec.get("mtime") or 0
    return (source_path or ("cloud:" + cloud_id), sheet, region, mtime)


def _region_png_path(key):
    """缓存 key → 缓存目录里的 PNG 路径。"""
    out_dir = os.path.join(tempfile.gettempdir(), "excel_send_wx_region")
    digest = hashlib.md5(repr(key).encode("utf-8")).hexdigest()[:16]
    return os.path.join(out_dir, "region_%s.png" % digest)


def ensure_region_png(spec, log_fn=None):
    """按渲染规格生成（并缓存）区域截图 PNG。

    spec 支持：
      source_path    本地 xlsx 路径
      cloud_file_id  云文档文件 ID（本地路径不可用时按需下载到临时文件）
      sheet          工作表名（空 = 第一个 sheet）
      range          "A1:F20"（空 = 已用区域）
      mtime          源文件修改时间（用于缓存失效）

    返回 PNG 路径；失败返回 None（调用方回退旧重绘逻辑）。
    """
    log = log_fn or (lambda _m: None)
    if not isinstance(spec, dict):
        return None

    source_path = str(spec.get("source_path") or "").strip()
    cloud_id = str(spec.get("cloud_file_id") or "").strip()
    sheet = str(spec.get("sheet") or "").strip()
    region = str(spec.get("range") or "").strip()

    key = _spec_cache_key(spec)
    cached = _PNG_CACHE.get(key)
    if cached and os.path.exists(cached):
        return cached

    out_png = _region_png_path(key)

    temp_path = None
    try:
        if not (source_path and os.path.isfile(source_path)):
            if not cloud_id:
                return None
            from .wps_cloud import (
                WpsOAuthStore, WpsCloudClient, cleanup_temp_file,
                new_temp_xlsx_path,
            )
            temp_path = new_temp_xlsx_path()
            log("区域截图：正在下载云文档以还原原始样式…")
            WpsCloudClient(WpsOAuthStore(), log_fn=log).download_to(
                cloud_id, temp_path
            )
            source_path = temp_path

        png = render_region_image(source_path, sheet, region, out_png, log_fn=log)
        if png:
            _PNG_CACHE[key] = png
        return png
    except Exception as exc:
        logger.exception("区域截图渲染失败")
        log(f"⚠ 区域截图渲染失败: {exc}")
        return None
    finally:
        if temp_path:
            try:
                from .wps_cloud import cleanup_temp_file
                cleanup_temp_file(temp_path)
            except Exception:
                pass


def _region_cache_file(path):
    """path 是缓存目录里的 region_*.png 时返回绝对路径，否则返回 None。"""
    if not path:
        return None
    try:
        abs_path = os.path.abspath(str(path))
    except Exception:
        return None
    cache_dir = os.path.abspath(
        os.path.join(tempfile.gettempdir(), "excel_send_wx_region")
    )
    name = os.path.basename(abs_path)
    if (os.path.dirname(abs_path).lower() != cache_dir.lower()
            or not name.startswith("region_") or not name.endswith(".png")):
        return None
    return abs_path


def is_cached_image(path):
    """是否为本模块渲染的缓存源图（整批发送可共用同一张，不按人删除）。"""
    return _region_cache_file(path) is not None


def drop_cached_file(path):
    """删除单张区域截图渲染缓存文件（只认本模块缓存目录里的 region_*.png）。

    用于一批发送结束后立即删除该批用过的源图，避免临时文件堆积；
    绝不误删用户文件 / 微信发送副本。返回是否删除成功。
    """
    abs_path = _region_cache_file(path)
    if not abs_path:
        return False

    removed = False
    try:
        if os.path.exists(abs_path):
            os.remove(abs_path)
            removed = True
    except OSError:
        removed = False
    for key, cached in list(_PNG_CACHE.items()):
        try:
            if cached and os.path.abspath(cached) == abs_path:
                _PNG_CACHE.pop(key, None)
        except Exception:
            pass
    return removed


def drop_cache_for(spec):
    """按渲染规格删除对应的缓存源图（一批发送结束后调用）。"""
    if not isinstance(spec, dict):
        return False
    return drop_cached_file(_region_png_path(_spec_cache_key(spec)))


def cleanup_cache():
    """清理区域截图渲染缓存（进程缓存 + 临时目录文件）。

    发送用的图片副本在每次发送结束后由 sender 删除；这里清理的是「按原表样式
    渲染」的源图缓存，避免临时文件跨会话堆积（应用启动时调用）。
    返回删除的文件数。
    """
    removed = 0
    for png in list(_PNG_CACHE.values()):
        try:
            if png and os.path.exists(png):
                os.remove(png)
                removed += 1
        except OSError:
            pass
    _PNG_CACHE.clear()

    folder = os.path.join(tempfile.gettempdir(), "excel_send_wx_region")
    if os.path.isdir(folder):
        for name in os.listdir(folder):
            if not name.startswith("region_") or not name.endswith(".png"):
                continue
            try:
                os.remove(os.path.join(folder, name))
                removed += 1
            except OSError:
                pass
    return removed
